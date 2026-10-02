#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GENERAR cosecha.json
====================

Convierte el Excel FLUJO_CEREZOS a un archivo liviano que la página de
Departamento Técnico puede leer sin internet.

    python3 scripts/generar_cosecha.py FLUJO_CEREZOS_2026-2027.xlsx

Toma tres hojas:
    FLUJO COSECHA 2026   el calendario de cosecha, día a día
    ESTIMACIONES         los momentos de estimación por cuartel
    KILOS SEMANA         el resumen de variedad por semana

Al final imprime un control de calidad: totales, nombres repetidos y
fechas que no calzan, para detectar errores del Excel antes de publicar.
"""

import datetime
import json
import os
import re
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl:  pip install openpyxl")
    sys.exit(2)


# ── utilidades ───────────────────────────────────────────────
def limpiar(v):
    """Texto sin espacios de más. Los nombres del Excel vienen con
    espacios al final que parten una variedad en dos."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return ""


def norm(s):
    s = unicodedata.normalize("NFD", limpiar(s).lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def semana_iso(fecha_iso):
    if not fecha_iso:
        return None
    y, m, d = map(int, fecha_iso.split("-"))
    return datetime.date(y, m, d).isocalendar()[1]


# ── hoja 1 · calendario de cosecha ───────────────────────────
def leer_flujo(wb, avisos):
    ws = wb["FLUJO COSECHA 2026"]
    filas_raw = list(ws.iter_rows(values_only=True))
    hdr = list(filas_raw[0])

    dias = [(i, iso(h)) for i, h in enumerate(hdr) if isinstance(h, datetime.datetime)]
    COL = {"prod": 1, "csg": 2, "rsocial": 3, "comuna": 4, "vari": 5,
           "momento": 6, "kg": 7, "ini": 9, "fin": 10, "emb": 11}

    out = []
    for row in filas_raw[1:]:
        if not row[COL["prod"]]:
            continue
        rec = {
            "cod": limpiar(row[COL["prod"]]),
            "csg": limpiar(row[COL["csg"]]),
            "prod": limpiar(row[COL["rsocial"]]).upper(),
            "comuna": limpiar(row[COL["comuna"]]).upper(),
            "vari": limpiar(row[COL["vari"]]).upper(),
            "momento": limpiar(row[COL["momento"]]),
            "kg": num(row[COL["kg"]]) or 0,
            "ini": iso(row[COL["ini"]]),
            "fin": iso(row[COL["fin"]]),
            "emb": num(row[COL["emb"]]),
        }
        # kilos por día, solo los que tienen valor
        d = {}
        for i, f in dias:
            v = num(row[i]) if i < len(row) else None
            if v:
                d[f] = v
        rec["dias"] = d
        out.append(rec)

    # ── control de calidad ──
    for r in out:
        if r["ini"] and r["fin"] and r["fin"] < r["ini"]:
            avisos.append("flujo: %s · %s termina (%s) antes de empezar (%s)"
                          % (r["prod"][:28], r["vari"], r["fin"], r["ini"]))
        s = sum(r["dias"].values())
        if r["kg"] and s and abs(s - r["kg"]) > max(1, r["kg"] * 0.01):
            avisos.append("flujo: %s · %s reparte %s kg en el calendario pero tiene %s comprometidos"
                          % (r["prod"][:28], r["vari"], f"{s:,.0f}".replace(",", "."),
                             f"{r['kg']:,.0f}".replace(",", ".")))
    return out, [f for _, f in dias]


# ── hoja 2 · estimaciones ────────────────────────────────────
MOMENTOS = ["Preliminar", "1° Estimacion", "2° Estimacion", "3° Estimacion",
            "4° Estimacion", "5° Estimacion", "6° Estimacion", "7° Estimacion",
            "8° Estimacion", "Estimacion Visual", "Estimacion Final"]

CALIBRES = ["4J", "3J", "2J", "J", "XL", "L"]


def leer_estimaciones(wb, avisos):
    ws = wb["ESTIMACIONES"]
    filas = list(ws.iter_rows(values_only=True))
    hdr = [limpiar(h).replace("\n", " ") for h in filas[0]]
    idx = {}
    for i, h in enumerate(hdr):
        if h and h not in idx:
            idx[h] = i

    def g(row, nombre):
        i = idx.get(nombre)
        return row[i] if (i is not None and i < len(row)) else None

    out = []
    for row in filas[1:]:
        if not row[1]:
            continue
        rec = {
            "cod": limpiar(row[1]),
            "csg": limpiar(row[2]),
            "prod": limpiar(row[3]).upper(),
            "comuna": limpiar(row[4]).upper(),
            "vari": limpiar(row[5]).upper(),
            "ha": num(row[6]),
            "cob": limpiar(row[7]),
            "usar": limpiar(g(row, "Estimacion a ocupar")),
            "embMenorJJ": num(g(row, "% EMBALAJE <JJ")),
            "descarte": num(g(row, "%descarte")),
            "embalaje": num(g(row, "% Embalaje")),
            "kgEmb": num(g(row, "Kg Embalados")),
            "pct2JUp": num(g(row, "% 2J Up")),
            "kg2JUp": num(g(row, "Kg 2J Up")),
            "kgHa2JUp": num(g(row, "Kg/Ha 2J Up")),
            "cosecha": limpiar(g(row, "Sistema de cosecha")),
            "clasif": limpiar(g(row, "CASIFICACION")),
            "nota": num(g(row, "PROMEDIO PONDERADO")),
            "vigor": num(g(row, "Nota de vigor 1=M,2=R,3=B,4=MB")),
        }
        rec["est"] = {m: (num(g(row, m)) or 0) for m in MOMENTOS}
        rec["cal"] = {c: num(g(row, c)) for c in CALIBRES}
        out.append(rec)

    # nombres escritos de dos formas
    for campo, etiqueta in (("vari", "variedad"), ("prod", "razón social")):
        grupos = {}
        for r in out:
            grupos.setdefault(norm(r[campo]), set()).add(r[campo])
        for k, v in grupos.items():
            if len(v) > 1:
                avisos.append("estimaciones: %s escrita de varias formas: %s"
                              % (etiqueta, " / ".join(sorted(v))))
    return out


# ── hoja 3 · kilos por semana ────────────────────────────────
def leer_kilos_semana(wb, avisos):
    ws = wb["KILOS SEMANA"]
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    # la cabecera es la fila que dice VARIEDAD / SEMANA
    fh = next((i for i, r in enumerate(filas)
               if any(limpiar(c).upper().startswith("VARIEDAD") for c in r if c)), None)
    if fh is None:
        avisos.append("kilos semana: no se encontró la fila de encabezado")
        return {"semanas": [], "filas": []}

    hdr = filas[fh]
    c0 = next(i for i, c in enumerate(hdr) if limpiar(c).upper().startswith("VARIEDAD"))
    semanas = []
    for i in range(c0 + 1, len(hdr)):
        v = hdr[i]
        # el Excel las guarda como texto ('42'), no como número
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            semanas.append((i, int(v)))
        elif isinstance(v, str) and v.strip().isdigit():
            semanas.append((i, int(v.strip())))

    out = []
    for row in filas[fh + 1:]:
        nombre = limpiar(row[c0]) if c0 < len(row) else ""
        if not nombre:
            continue
        if norm(nombre).startswith("total"):
            continue           # el total se recalcula, no se copia
        vals = {}
        for i, s in semanas:
            v = num(row[i]) if i < len(row) else None
            if v:
                vals[str(s)] = v
        out.append({"vari": nombre.upper(), "kg": vals})

    # variedades repetidas por diferencias de escritura
    grupos = {}
    for r in out:
        grupos.setdefault(norm(r["vari"]), []).append(r["vari"])
    for k, v in grupos.items():
        if len(v) > 1:
            avisos.append("kilos semana: %s aparece %d veces (espacios o mayúsculas distintas)"
                          % (v[0], len(v)))
    return {"semanas": [s for _, s in semanas], "filas": out}


# ── principal ────────────────────────────────────────────────
def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        print("No existe el archivo:", ruta)
        return 2

    print("Leyendo", os.path.basename(ruta), "...")
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    avisos = []

    flujo, dias = leer_flujo(wb, avisos)
    estim = leer_estimaciones(wb, avisos)
    kilos = leer_kilos_semana(wb, avisos)
    wb.close()

    # semanas del calendario, derivadas de las fechas
    semanas = {}
    for f in dias:
        semanas.setdefault(semana_iso(f), []).append(f)

    # Kilos Semana se recalcula desde el calendario en vez de copiar la
    # tabla dinámica del Excel: esa se queda atrás cuando se agrega un
    # productor y nadie la refresca. Así siempre cuadra.
    calc = {}
    for r in flujo:
        v = r["vari"]
        for f, kg in r["dias"].items():
            sem = semana_iso(f)
            calc.setdefault(v, {}).setdefault(str(sem), 0)
            calc[v][str(sem)] += kg
    semanas_calc = sorted({int(s) for v in calc.values() for s in v},
                          key=lambda n: (n < 40, n))
    kilos_calc = {
        "semanas": semanas_calc,
        "filas": [{"vari": v, "kg": {k: round(x, 2) for k, x in sorted(d2.items())}}
                  for v, d2 in sorted(calc.items())],
    }

    paquete = {
        "meta": {
            "generado": datetime.date.today().isoformat(),
            "origen": os.path.basename(ruta),
            "temporada": "2026-2027",
            "dias": dias,
            "semanas": [{"n": n, "dias": ds} for n, ds in sorted(semanas.items(),
                        key=lambda x: x[1][0])],
        },
        "flujo": flujo,
        "estimaciones": estim,
        "kilosSemana": kilos_calc,
        "kilosSemanaExcel": kilos,
    }

    salida = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "cosecha.json")
    with open(salida, "w", encoding="utf-8") as f:
        json.dump(paquete, f, ensure_ascii=False, separators=(",", ":"))

    mil = lambda n: f"{n:,.0f}".replace(",", ".")
    kg_flujo = sum(r["kg"] for r in flujo)
    kg_dias = sum(sum(r["dias"].values()) for r in flujo)
    kg_sem = sum(sum(r["kg"].values()) for r in kilos["filas"])
    kg_calc = sum(sum(r["kg"].values()) for r in kilos_calc["filas"])

    print("\nRESUMEN")
    print("  calendario    : %d filas · %d productores · %d variedades"
          % (len(flujo), len(set(r["prod"] for r in flujo)),
             len(set(r["vari"] for r in flujo))))
    print("  estimaciones  : %d filas · %d productores" % (len(estim),
          len(set(r["prod"] for r in estim))))
    print("  kilos semana  : %d variedades × %d semanas"
          % (len(kilos["filas"]), len(kilos["semanas"])))
    print("  días del flujo: %s a %s (%d días, %d semanas)"
          % (dias[0] if dias else "—", dias[-1] if dias else "—",
             len(dias), len(semanas)))
    print("\nKILOS")
    print("  comprometidos en el calendario : %s" % mil(kg_flujo))
    print("  repartidos día a día           : %s" % mil(kg_dias))
    print("  Kilos Semana recalculado       : %s" % mil(kg_calc))
    print("  Kilos Semana del Excel         : %s" % mil(kg_sem))
    if abs(kg_calc - kg_dias) > 1:
        print("  ⚠ el recálculo no cuadra con el calendario: %s kg de diferencia"
              % mil(abs(kg_calc - kg_dias)))
    else:
        print("  ✓ el recálculo cuadra exacto con el calendario")
    if not kg_sem:
        print("  ⚠ no se leyó ningún kilo de la hoja Kilos Semana del Excel")
    elif abs(kg_calc - kg_sem) > 1:
        print("  ⚠ la tabla del Excel está %s kg atrás — conviene refrescarla allá"
              % mil(kg_calc - kg_sem))

    if avisos:
        print("\nREVISAR EN EL EXCEL (%d):" % len(avisos))
        for a in avisos[:20]:
            print("  ·", a)
        if len(avisos) > 20:
            print("  … y %d más" % (len(avisos) - 20))

    print("\ncosecha.json → %.0f KB" % (os.path.getsize(salida) / 1024))
    return 0


if __name__ == "__main__":
    sys.exit(main())
