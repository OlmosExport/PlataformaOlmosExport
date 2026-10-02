#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GENERAR LOS ARCHIVOS DE LA AGENDA DE PESTICIDAS
===============================================

Convierte la AGENDA DE PESTICIDAS de ASOEX en archivos que la sección
Requisitos SAG puede buscar sin internet y sin pedir login.

    python3 scripts/generar_pesticidas.py AGENDA_PESTICIDAS_2026.xlsx

Genera:

    pesticidas.json              índice liviano (especies, fechas, mercados)
    pesticidas_<especie>.json    una especie por archivo, se carga al abrirla

Cada hoja del Excel es una especie y trae SU PROPIA fecha de descarga,
que no siempre coincide con el nombre del archivo ni con las demás hojas.
El índice guarda la fecha de cada una para poder avisar cuál está vieja.

OJO — REPARACIÓN DE LMR CONVERTIDOS EN FECHA
--------------------------------------------
En el Excel de origen hay celdas de LMR que Excel transformó en fecha:

  · Un LMR simple con formato de fecha. "2" quedó como 1900-01-02.
    Se recupera volviendo al número de serie de Excel.

  · Un par de LMR escrito "2/4" (una sustancia combinada, un valor para
    cada una) que Excel leyó como el 2 de abril. Se recupera como
    día/mes, que es el orden en que fue escrito.

Que el orden es día/mes está comprobado contra las filas de sustancia
simple: CHILE AZOXYSTROBIN = 2 ppm, y el par CHILE
AZOXYSTROBIN/TEBUCONAZOLE quedó como 2026-04-02, o sea "2/4".
Lo mismo con CANADA (2026-03-02 → "2/3", azoxystrobin 2 ppm).

El script cuenta cuántas celdas reparó y lo informa al final.
"""

import datetime
import json
import os
import re
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl:  pip install openpyxl")
    sys.exit(2)

# Posición de cada dato en la fila (encabezado en la fila 7, datos desde la 8)
COL = {"especie": 2, "mercado": 3, "sustancia": 4, "tipo": 5,
       "ppm": 6, "ppm_obs": 7, "dias": 8, "dias_obs": 9,
       "nuevo_lmr": 10, "vigencia": 11}

# Un LMR que Excel guardó como fecha anterior a este día es un número suelto
# (el "1900-01-02" de Excel es el número 2). Después de ese día ya no puede
# ser un LMR: es un par "día/mes" que Excel malinterpretó.
CORTE_NUMERO = datetime.date(1900, 6, 30)


def quitar_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", quitar_tildes(s).lower()).strip("_")


def serie_excel(d):
    """Número que tenía la celda antes de que Excel la mostrara como fecha."""
    if d < datetime.date(1900, 3, 1):          # antes del bug del año 1900
        return (d - datetime.date(1899, 12, 31)).days
    return (d - datetime.date(1899, 12, 30)).days


def texto(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s in ("/", "-") else s


def limpiar_lmr(v, reparos):
    """
    Devuelve el LMR tal como debería leerse.
    reparos es una lista donde deja constancia de cada celda arreglada.
    """
    if isinstance(v, (datetime.datetime, datetime.date)):
        d = v.date() if isinstance(v, datetime.datetime) else v
        if d <= CORTE_NUMERO:
            arreglado = str(serie_excel(d))
            reparos.append(("numero", d.isoformat(), arreglado))
        else:
            arreglado = "%d/%d" % (d.day, d.month)
            reparos.append(("par", d.isoformat(), arreglado))
        return arreglado
    return texto(v)


def numero(s):
    """El LMR puede ser un número, un código (ST, EX) o un par 'a/b'."""
    if not s:
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def fecha_descarga(ws):
    """Cada hoja declara su fecha de descarga en las primeras filas."""
    for row in ws.iter_rows(min_row=1, max_row=7, max_col=8, values_only=True):
        for i, v in enumerate(row):
            if v and "fecha descarga" in str(v).lower():
                for w in row[i + 1:]:
                    f = texto(w)
                    if f:
                        return normalizar_fecha(f)
    return ""


def normalizar_fecha(f):
    """Las hojas traen 2026-01-05 o 15-11-2024. Dejar todo en ISO."""
    m = re.match(r"^(\d{2})-(\d{2})-(\d{4})$", f)
    if m:
        return "%s-%s-%s" % (m.group(3), m.group(2), m.group(1))
    return f


def nombre_especie(hoja):
    """'CEREZAS 05-01-2026' → 'CEREZAS'."""
    s = re.sub(r"[\d\-/_.\s]+$", "", hoja).strip()
    return (s or hoja).upper()


def leer_hoja(ws):
    filas, notas, reparos = [], [], []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not row or len(row) <= COL["tipo"] or not row[COL["especie"]]:
            continue
        g = lambda k: texto(row[COL[k]]) if len(row) > COL[k] else ""
        mercado, sustancia = g("mercado"), g("sustancia")
        if not mercado or not sustancia:
            # Al final de la hoja vienen las notas al pie y el significado
            # de las abreviaturas. Se guardan: son parte del documento.
            n = texto(row[COL["especie"]])
            if n:
                notas.append(n)
            continue
        ppm = limpiar_lmr(row[COL["ppm"]] if len(row) > COL["ppm"] else None, reparos)
        dias = limpiar_lmr(row[COL["dias"]] if len(row) > COL["dias"] else None, reparos)
        filas.append({
            "mercado": mercado, "sustancia": sustancia, "tipo": g("tipo"),
            "ppm": ppm, "ppmNum": numero(ppm), "obs": g("ppm_obs"),
            "dias": dias, "diasNum": numero(dias), "diasObs": g("dias_obs"),
            "nuevo": g("nuevo_lmr"), "vig": g("vigencia"),
        })
    return filas, notas, reparos


def comprimir(filas):
    """
    Guardar cada fila como un arreglo y los textos repetidos una sola vez.
    Un mercado como 'CODEX - CXL' se escribe 200 veces; así se escribe una.
    """
    mercados, sustancias, tipos = [], [], []
    im, isus, it = {}, {}, {}

    def idx(v, lista, mapa):
        if v not in mapa:
            mapa[v] = len(lista)
            lista.append(v)
        return mapa[v]

    datos = []
    for f in filas:
        datos.append([
            idx(f["mercado"], mercados, im),
            idx(f["sustancia"], sustancias, isus),
            idx(f["tipo"], tipos, it),
            f["ppm"], f["ppmNum"], f["obs"],
            f["dias"], f["diasNum"], f["diasObs"],
            f["nuevo"], f["vig"],
        ])
    return {
        "campos": ["mercado", "sustancia", "tipo", "ppm", "ppmNum", "obs",
                   "dias", "diasNum", "diasObs", "nuevo", "vig"],
        "mercados": mercados, "sustancias": sustancias, "tipos": tipos,
        "filas": datos,
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        print("No existe el archivo:", ruta)
        return 2

    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    print("Leyendo", os.path.basename(ruta), "...\n")
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    indice, avisos, total, reparados = [], [], 0, {"numero": 0, "par": 0}
    ejemplos = []

    for hoja in wb.sheetnames:
        ws = wb[hoja]
        esp = nombre_especie(hoja)
        fd = fecha_descarga(ws)
        filas, notas, reparos = leer_hoja(ws)

        if not filas:
            avisos.append("la hoja '%s' no trajo ninguna fila utilizable" % hoja)
            continue

        for clase, antes, despues in reparos:
            reparados[clase] += 1
            if len(ejemplos) < 4:
                ejemplos.append("%s · %s → %s" % (esp, antes, despues))

        paquete = comprimir(filas)
        paquete["especie"] = esp
        paquete["hoja"] = hoja
        paquete["fecha"] = fd
        paquete["notas"] = notas

        archivo = "pesticidas_%s.json" % slug(esp)
        with open(os.path.join(base, archivo), "w", encoding="utf-8") as f:
            json.dump(paquete, f, ensure_ascii=False, separators=(",", ":"))

        peso = os.path.getsize(os.path.join(base, archivo)) / 1024
        indice.append({
            "especie": esp, "archivo": archivo, "fecha": fd,
            "filas": len(filas),
            "mercados": len(paquete["mercados"]),
            "sustancias": len(paquete["sustancias"]),
            "cambios": sum(1 for x in filas if x["nuevo"]),
        })
        total += len(filas)
        print("  %-22s %5d filas · %2d mercados · %3d sustancias · "
              "%3d cambios · %s · %4.0f KB · %s"
              % (esp, len(filas), len(paquete["mercados"]),
                 len(paquete["sustancias"]),
                 indice[-1]["cambios"], fd or "sin fecha", peso,
                 "%d reparados" % len(reparos) if reparos else "-"))
    wb.close()

    # ── la fecha del nombre del archivo rara vez vale para todas las hojas ──
    nombre = os.path.basename(ruta)
    m = re.search(r"(\d{2})[-_](\d{2})[-_](\d{4})", nombre)
    f_nombre = ("%s-%s-%s" % (m.group(3), m.group(2), m.group(1))) if m else ""
    fechas = sorted({i["fecha"] for i in indice if i["fecha"]})
    hoy = datetime.date.today()

    for i in indice:
        if not i["fecha"]:
            avisos.append("%s no declara fecha de descarga" % i["especie"])
            continue
        try:
            dias = (hoy - datetime.date.fromisoformat(i["fecha"])).days
        except ValueError:
            continue
        i["dias"] = dias
        if dias > 180:
            avisos.append("%s se descargó hace %d días (%s)"
                          % (i["especie"], dias, i["fecha"]))
    if len(fechas) > 1:
        avisos.append("las hojas no se descargaron el mismo día: %s"
                      % ", ".join(fechas))
    if f_nombre and f_nombre not in fechas:
        avisos.append("el archivo se llama %s y ninguna hoja tiene esa fecha"
                      % f_nombre)

    idx = {
        "meta": {
            "generado": hoy.isoformat(),
            "origen": nombre,
            "fechaNombre": f_nombre,
            "fechaMasReciente": fechas[-1] if fechas else "",
            "fechaMasAntigua": fechas[0] if fechas else "",
            "registros": total,
            "fuente": "Agenda de Pesticidas ASOEX",
        },
        "especies": indice,
    }
    with open(os.path.join(base, "pesticidas.json"), "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False, indent=1)

    print("\nRESUMEN")
    print("  especies            : %d" % len(indice))
    print("  registros           : %s" % f"{total:,}".replace(",", "."))
    print("  fechas de descarga  : %s" % ", ".join(fechas))
    print("  índice              : %.0f KB"
          % (os.path.getsize(os.path.join(base, "pesticidas.json")) / 1024))

    if reparados["numero"] or reparados["par"]:
        print("\nLMR REPARADOS (el Excel los tenía convertidos en fecha)")
        print("  número suelto       : %d" % reparados["numero"])
        print("  par 'a/b'           : %d" % reparados["par"])
        for e in ejemplos:
            print("    ", e)

    if avisos:
        print("\nREVISAR")
        for a in avisos:
            print("  ·", a)
    return 0


if __name__ == "__main__":
    sys.exit(main())
