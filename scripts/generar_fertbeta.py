#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Lee el libro 'Programa de Fertilizacion Agricola Antumapu 26-27' (formato
unificado Abud & Cia. v6.0) y lo deja como fertbeta.json para la plataforma.

Del libro salen seis cosas:

  catalogo       hoja 'Fertilizantes'  - 31 productos con sus leyes (%)
  listas         hoja 'Listas'         - especies, etapas, vigor, grupos, portainjertos
  fenologia      hoja 'Fenologia'      - 24 combinaciones especie+grupo con sus dias
  inicios        hoja 'Fenologia'      - dia y mes de inicio por especie+grupo, y el ano
  requerimientos hoja 'Requerimientos' - factor a oxido y factores de absorcion
  ejemplos       hojas 'Hayward ...'   - los 4 programas reales, como punto de partida

No se recalcula nada aqui: el motor vive en el navegador. Esto solo copia los
datos del libro para que la pagina los tenga sin depender del Excel.

Uso:  python3 generar_fertbeta.py <libro.xlsx> [salida.json]
"""

import json
import sys
from pathlib import Path

import openpyxl

# El libro reserva Fertilizantes!B6:B60 para el catalogo y Fenologia!A8:A304
# para los estados. Respetamos esos limites.
CAT_D, CAT_H = 6, 60
FEN_D, FEN_H = 8, 304
INI_D, INI_H = 8, 31
LISTA_D, LISTA_H = 4, 40

# Columnas de la hoja de programa (MODELO y las copias).
COL_BASE = list(range(5, 12))    # E..K  - 7 fertilizantes base
COL_COMP = list(range(13, 16))   # M..O  - 3 complementarios
FILA_D, FILA_H = 10, 50          # 41 semanas
FILA_TOT = 51

NUTRIENTES = ["N", "P2O5", "K2O", "CaO", "MgO"]


def txt(v):
    return "" if v is None else str(v).strip()


def num(v):
    if v is None or isinstance(v, str):
        return None
    try:
        return round(float(v), 6)
    except (TypeError, ValueError):
        return None


def columna(hoja, col, desde, hasta):
    """Valores no vacios de una columna, en orden."""
    out = []
    for f in range(desde, hasta + 1):
        v = txt(hoja.cell(f, col).value)
        if v:
            out.append(v)
    return out


def leer_catalogo(hoja):
    """Hoja 'Fertilizantes': un producto por fila, con sus leyes."""
    prods = []
    for f in range(CAT_D, CAT_H + 1):
        nombre = txt(hoja.cell(f, 2).value)
        codigo = txt(hoja.cell(f, 1).value)
        # El libro deja notas al pie dentro del rango de busqueda (filas 38-40,
        # sin codigo). Un producto de verdad siempre trae codigo.
        if not nombre or not codigo:
            continue
        prods.append({
            "codigo": codigo,
            "nombre": nombre,
            "tipo": txt(hoja.cell(f, 3).value),          # Base / Complementario
            "categoria": txt(hoja.cell(f, 4).value),
            "N": num(hoja.cell(f, 5).value) or 0,
            "P2O5": num(hoja.cell(f, 6).value) or 0,
            "K2O": num(hoja.cell(f, 7).value) or 0,
            "CaO": num(hoja.cell(f, 8).value) or 0,
            "MgO": num(hoja.cell(f, 9).value) or 0,
            "S": num(hoja.cell(f, 10).value) or 0,
            "unidad": txt(hoja.cell(f, 11).value) or "kg",
            "obs": txt(hoja.cell(f, 12).value),
        })
    return prods


def leer_listas(hoja):
    """Hoja 'Listas': las listas desplegables, incluidas las que se filtran."""
    especies = columna(hoja, 3, LISTA_D, LISTA_H)
    # Las columnas G..K son GRUPO_<especie> y M..Q son PORTA_<especie>, en el
    # mismo orden en que el libro declara los nombres definidos.
    orden_grupo = ["Cerezo", "Kiwi", "Ciruelo", "Manzano", "Peral"]
    grupos = {e: columna(hoja, 7 + i, LISTA_D, LISTA_H)
              for i, e in enumerate(orden_grupo)}
    portas = {e: columna(hoja, 13 + i, LISTA_D, LISTA_H)
              for i, e in enumerate(orden_grupo)}
    return {
        "especies": especies,
        "etapas": columna(hoja, 4, LISTA_D, LISTA_H),
        "vigor": columna(hoja, 5, LISTA_D, LISTA_H),
        "grupos": grupos,
        "portainjertos": portas,
        "base": columna(hoja, 1, LISTA_D, LISTA_H),
        "complementarios": columna(hoja, 2, LISTA_D, LISTA_H),
    }


def leer_fenologia(hoja):
    """Hoja 'Fenologia': estados y dias por especie+grupo, y las fechas de inicio."""
    combos = {}
    for f in range(FEN_D, FEN_H + 1):
        esp = txt(hoja.cell(f, 2).value)
        grupo = txt(hoja.cell(f, 3).value)
        if not esp or not grupo:
            continue
        orden = num(hoja.cell(f, 4).value)
        estado = txt(hoja.cell(f, 5).value)
        dias = num(hoja.cell(f, 6).value)
        if orden is None or dias is None or not estado:
            continue
        combos.setdefault(esp + "|" + grupo, []).append({
            "orden": int(orden), "estado": estado, "dias": int(dias),
        })
    for v in combos.values():
        v.sort(key=lambda x: x["orden"])

    inicios = {}
    for f in range(INI_D, INI_H + 1):
        esp = txt(hoja.cell(f, 9).value)
        grupo = txt(hoja.cell(f, 10).value)
        if not esp or not grupo:
            continue
        dia = num(hoja.cell(f, 11).value)
        mes = num(hoja.cell(f, 12).value)
        if dia is None or mes is None:
            continue
        inicios[esp + "|" + grupo] = {"dia": int(dia), "mes": int(mes)}

    ano = num(hoja.cell(3, 11).value)      # K3
    return combos, inicios, int(ano) if ano else None


def leer_requerimientos(hoja):
    """Hoja 'Requerimientos': conversion a oxido y factores de absorcion."""
    oxido, absorcion = {}, {}
    for i, f in enumerate(range(6, 11)):
        elem = txt(hoja.cell(f, 1).value)
        if elem:
            oxido[elem] = num(hoja.cell(f, 3).value)
        nut = txt(hoja.cell(f, 5).value)
        if nut:
            absorcion[nut] = {
                "Cerezo": num(hoja.cell(f, 6).value),
                "Kiwi": num(hoja.cell(f, 7).value),
                "Otras": num(hoja.cell(f, 8).value),
            }
    # Valores referenciales de analisis de fruta (mg/100 g).
    refs = []
    for f in range(17, 22):
        esp = txt(hoja.cell(f, 1).value)
        if not esp:
            continue
        vals = [num(hoja.cell(f, c).value) for c in range(3, 8)]
        if all(v is None for v in vals):
            continue
        refs.append({
            "especie": esp,
            "referencia": txt(hoja.cell(f, 2).value),
            "N": vals[0], "P": vals[1], "K": vals[2],
            "Ca": vals[3], "Mg": vals[4],
        })
    return {"oxido": oxido, "absorcion": absorcion, "referencias": refs}


def leer_programa(hoja, nombre):
    """Una hoja de programa: encabezado, columnas elegidas y dosis por semana."""
    base = [txt(hoja.cell(9, c).value) for c in COL_BASE]
    comp = [txt(hoja.cell(9, c).value) for c in COL_COMP]

    dosis_base, dosis_comp = {}, {}
    for f in range(FILA_D, FILA_H + 1):
        i = f - FILA_D
        for j, c in enumerate(COL_BASE):
            v = num(hoja.cell(f, c).value)
            if v is not None:
                dosis_base[f"{i},{j}"] = v
        for j, c in enumerate(COL_COMP):
            v = num(hoja.cell(f, c).value)
            if v is not None:
                dosis_comp[f"{i},{j}"] = v

    fruta = {}
    for k, f in zip(["N", "P", "K", "Ca", "Mg"], range(63, 68)):
        v = num(hoja.cell(f, 4).value)
        if v is not None:
            fruta[k] = v

    return {
        "nombre": nombre,
        "productor": txt(hoja.cell(3, 10).value),
        "campo": txt(hoja.cell(3, 16).value),
        "especie": txt(hoja.cell(4, 4).value),
        "variedad": txt(hoja.cell(4, 10).value),
        "grupo": txt(hoja.cell(4, 16).value),
        "porta": txt(hoja.cell(5, 4).value),
        "etapa": txt(hoja.cell(5, 10).value),
        "vigor": txt(hoja.cell(5, 16).value),
        "ha": num(hoja.cell(6, 4).value),
        "ton": num(hoja.cell(6, 10).value),
        "base": base,
        "comp": comp,
        "dosis": dosis_base,
        "dosisComp": dosis_comp,
        "fruta": fruta,
    }


def leer_leeme(hoja):
    """Las instrucciones del libro, para mostrarlas en la pagina de ayuda."""
    out = []
    for f in range(2, 73):
        v = txt(hoja.cell(f, 2).value)
        out.append(v)
    # Quitamos las lineas de cola vacias.
    while out and not out[-1]:
        out.pop()
    return out


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    origen = Path(sys.argv[1])
    destino = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("fertbeta.json")

    w = openpyxl.load_workbook(origen, data_only=False)

    catalogo = leer_catalogo(w["Fertilizantes"])
    listas = leer_listas(w["Listas"])
    fenologia, inicios, ano = leer_fenologia(w["Fenologia"])
    requerimientos = leer_requerimientos(w["Requerimientos"])

    ejemplos = [leer_programa(w[s], s) for s in w.sheetnames
                if s.startswith("Hayward")]

    # Avisos: lo que el libro declara en una hoja y falta en otra.
    avisos = []
    nombres_cat = {p["nombre"] for p in catalogo}
    for grupo, etiqueta in (("base", "FERTILIZANTES BASE"),
                            ("complementarios", "COMPLEMENTARIOS")):
        for n in listas[grupo]:
            if n not in nombres_cat:
                avisos.append(
                    f"«{n}» esta en la lista {etiqueta} pero no en el catalogo "
                    f"de fertilizantes, asi que sus leyes entran como cero.")
    for p in ejemplos:
        for n in p["base"] + p["comp"]:
            if n and n not in nombres_cat:
                avisos.append(
                    f"El programa «{p['nombre']}» usa «{n}», que no esta en el "
                    f"catalogo de fertilizantes.")
        for n in p["base"]:
            if n and n in nombres_cat:
                tipo = next(x["tipo"] for x in catalogo if x["nombre"] == n)
                if tipo == "Complementario":
                    avisos.append(
                        f"El programa «{p['nombre']}» lleva «{n}» en una columna "
                        f"de fertilizantes base, pero el catalogo lo declara "
                        f"complementario: suma al balance de unidades.")

    datos = {
        "fuente": origen.name,
        "formato": "Abud & Cia. - formato unificado v6.0",
        "temporada": "2026-2027",
        "anoTemporada": ano,
        "catalogo": catalogo,
        "listas": listas,
        "fenologia": fenologia,
        "inicios": inicios,
        "requerimientos": requerimientos,
        "ejemplos": ejemplos,
        "leeme": leer_leeme(w["LEEME"]),
        "avisos": sorted(set(avisos)),
    }

    destino.write_text(json.dumps(datos, ensure_ascii=False, indent=1),
                       encoding="utf-8")

    print(f"{destino}  ({destino.stat().st_size/1024:.1f} KB)")
    print(f"  catalogo        {len(catalogo)} productos "
          f"({sum(1 for p in catalogo if p['tipo']=='Base')} base, "
          f"{sum(1 for p in catalogo if p['tipo']=='Complementario')} complementarios)")
    print(f"  fenologia       {len(fenologia)} combinaciones especie+grupo")
    print(f"  inicios         {len(inicios)} fechas de inicio, ano {ano}")
    print(f"  listas          {len(listas['especies'])} especies, "
          f"{len(listas['base'])} base, {len(listas['complementarios'])} complementarios")
    print(f"  ejemplos        {len(ejemplos)} programas: "
          + ", ".join(p["nombre"] for p in ejemplos))
    if avisos:
        print(f"  avisos          {len(set(avisos))}")
        for a in sorted(set(avisos)):
            print(f"     - {a}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
